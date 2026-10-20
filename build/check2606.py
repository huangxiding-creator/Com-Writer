# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
doc = Document(r"e:\AI-Station\Com-Writer\02-1 总承包事业部\05 成果迭代优化\20260622 黑河黄藏寺水利枢纽工程专题会会议纪要（定稿）0624.docx")
print("段落数:", len(doc.paragraphs), "表格数:", len(doc.tables))
for t in doc.tables:
    for row in t.rows:
        print(" | ".join(c.text.strip().replace("\n","⏎")[:80] for c in row.cells))
# textbox 检查
import re
xml = doc.element.xml
texts = re.findall(r'<w:t[^>]*>([^<]{4,})</w:t>', xml)
print("--- 前12条长文本 ---")
for t in texts[:12]:
    print(t[:70])
