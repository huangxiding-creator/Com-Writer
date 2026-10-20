# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
doc = Document(r"e:\AI-Station\Com-Writer\02-2 黄藏寺EPC项目\13 AI输出初稿\黑河黄藏寺水利枢纽工程专题会会议纪要（初稿）260809.docx")
for p in doc.paragraphs:
    t = p.text.strip()
    if t:
        print(f"[{p.style.name}] {t}")
for tb in doc.tables:
    print(f"[TABLE {len(tb.rows)}x{len(tb.columns)}]")
