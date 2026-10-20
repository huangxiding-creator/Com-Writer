# -*- coding: utf-8 -*-
"""黄藏寺纪要生成器 v2 —— 模板A底版原位改写（哈主任十二律配套）。

用法:
  E:/AI-Station/.venv/Scripts/python.exe gen_minutes_v2.py <内容JSON> <输出docx>

内容JSON格式:
{
  "meeting_no": "〔2026〕XX号",            # 保持占位
  "topic": "黑河黄藏寺水利枢纽工程专题会",   # 议题(不加括号扩展)
  "date": "2026年10月12日上午",
  "place": "211会议室",
  "host": "哈佳",
  "attendees": "哈佳、齐三红、于立新、张艳锋、贾 超、王 飞（线上）、黄细丁",
  "opening": "2026年10月12日上午，总包部主任哈佳主持召开了……形成纪要如下：",
  "items": ["一、……", "二、……", ...],     # 8-15条, 每条已按十二律写成
  "this_page_no_text": true                 # 末页发送栏独页时加（本页无正文）
}
生成后字体统一仿宋_GB2312 16pt、首行缩进38pt、行距29pt（模板继承+兜底注入）。
"""
import copy
import sys, io, os, json
if sys.stdout.encoding != 'utf-8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

TEMPLATE = r"e:\AI-Station\Com-Writer\02-2 黄藏寺EPC项目\总包部黑河黄藏寺水利枢纽工程专题会（模板A）.docx"


def set_para_text(p, text):
    """整段替换文本；rPr优先取带eastAsia字体的run；兜底注入仿宋_GB2312。"""
    ref, ref_fallback = None, None
    for r in p.runs:
        rPr_el = r._element.find(qn('w:rPr'))
        if rPr_el is None:
            continue
        if ref_fallback is None:
            ref_fallback = rPr_el
        rf = rPr_el.find(qn('w:rFonts'))
        if rf is not None and rf.get(qn('w:eastAsia')):
            ref = rPr_el
            break
    if ref is None:
        ref = ref_fallback
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    r_el = run._element
    old = r_el.find(qn('w:rPr'))
    new_rPr = copy.deepcopy(ref) if ref is not None else old
    if old is not None:
        r_el.remove(old)
    if new_rPr is None:
        new_rPr = OxmlElement('w:rPr')
    rf = new_rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        new_rPr.insert(0, rf)
    rf.set(qn('w:eastAsia'), '仿宋_GB2312')
    rf.set(qn('w:ascii'), '仿宋_GB2312')
    rf.set(qn('w:hAnsi'), '仿宋_GB2312')
    r_el.insert(0, new_rPr)
    return p


def gen(content, dst):
    doc = Document(TEMPLATE)
    ps = doc.paragraphs
    # 模板段位: 0标题 1文号 4议题 5时间 6地点 7主持 8参会 9纪要内容: 10导语 11..16条目 …25分隔 26发送 28整理人
    set_para_text(ps[4], f"会议议题：{content['topic']}")
    set_para_text(ps[5], f"会议时间：{content['date']}")
    set_para_text(ps[6], f"会议地点：{content['place']}")
    set_para_text(ps[7], f"主 持 人：{content['host']}")
    set_para_text(ps[8], f"参加人员：{content['attendees']}")
    set_para_text(ps[10], content['opening'])

    proto = ps[11]._element
    anchor = None
    for i, text in enumerate(content['items']):
        idx = 11 + i
        if idx <= 16:
            set_para_text(ps[idx], text)
            anchor = ps[idx]._element
        else:
            new_el = copy.deepcopy(proto)
            anchor.addnext(new_el)
            anchor = new_el
            set_para_text(Paragraph(new_el, ps[11]._parent), text)

    # 模板原有条目多于本次：删除多余条目段（11+len .. 16）
    for idx in range(11 + len(content['items']), 17):
        if idx < len(ps):
            el = ps[idx]._element
            el.getparent().remove(el)

    # （本页无正文）：可选，插在首条分隔线段之前
    if content.get('this_page_no_text'):
        for p in doc.paragraphs:
            if p.text.strip() == "____________________________________________________________":
                new_el = copy.deepcopy(ps[10]._element)
                p._element.addprevious(new_el)
                set_para_text(Paragraph(new_el, ps[10]._parent), "（本页无正文）")
                break

    doc.save(dst)
    print("saved:", dst)

    # 结构回读
    d2 = Document(dst)
    marks = tuple(f"{c}、" for c in "一二三四五六七八九十") + ("十一、","十二、","十三、","十四、","十五、")
    n = sum(1 for p in d2.paragraphs
            if p.text.strip().startswith(marks))
    total = sum(len(p.text) for p in d2.paragraphs)
    print(f"条目数={n} 全文字符={total}")


if __name__ == "__main__":
    content_path, dst = sys.argv[1], sys.argv[2]
    with open(content_path, encoding='utf-8') as f:
        content = json.load(f)
    gen(content, dst)
