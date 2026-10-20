# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn

DST = r"e:\AI-Station\Com-Writer\02-2 黄藏寺EPC项目\13 AI输出初稿\黑河黄藏寺水利枢纽工程专题会会议纪要（初稿）261012.docx"
doc = Document(DST)
issues = []
n_item = 0
for p in doc.paragraphs:
    t = p.text.strip()
    if not t:
        continue
    pf = p.paragraph_format
    # 条目段格式核查
    for mark in ("一、","二、","三、","四、","五、","六、","七、","八、","九、","十、","十一、","十二、","十三、","十四、"):
        if t.startswith(mark):
            n_item += 1
            fi = pf.first_line_indent
            if fi is None or abs(fi.pt - 38) > 2:
                issues.append(f"条目{mark} 首行缩进异常: {fi}")
            ls = pf.line_spacing
            if ls is None or abs(ls.pt - 29) > 1.5:
                issues.append(f"条目{mark} 行距异常: {ls}")
            for r in p.runs:
                if r.text.strip():
                    ea = None
                    rPr = r._element.find(qn('w:rPr'))
                    if rPr is not None:
                        rf = rPr.find(qn('w:rFonts'))
                        if rf is not None:
                            ea = rf.get(qn('w:eastAsia'))
                    if ea != '仿宋_GB2312':
                        issues.append(f"条目{mark} 字体异常: {ea}")
                    if r.font.size is None or abs(r.font.size.pt - 16) > 0.6:
                        issues.append(f"条目{mark} 字号异常: {r.font.size}")
            break
print(f"条目总数: {n_item}")
print("格式问题:", issues if issues else "无")

# 字数统计
total = sum(len(p.text) for p in doc.paragraphs)
print(f"全文字符数: {total}")
